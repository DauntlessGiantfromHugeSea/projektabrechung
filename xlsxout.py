"""
Excel-Erzeugung (.xlsx) fuer Berichte/Buchungslisten.
Liefert reine Bytes, die als Download bereitgestellt oder gespeichert werden.
"""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font, PatternFill

import config
from events import Interval


def _fmt_dur(hours: float) -> str:
    m = round(hours * 60)
    return f"{m // 60}:{m % 60:02d}"


def intervals_xlsx(intervals: list[Interval], title: str = "Bericht") -> bytes:
    """Buchungsliste (Detail) + Summen je Mitarbeiter als XLSX."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Buchungen"
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF5E9")

    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Datum", "Mitarbeiter", "Projekt", "Kommt", "Geht",
               "Dauer (Std:Min)", "Stunden (dez.)", "Tätigkeit", "Quelle"])
    for c in ws[3]:
        c.font = head
        c.fill = fill

    total = 0.0
    per_emp: dict[str, float] = {}
    for iv in sorted(intervals, key=lambda x: x.start):
        st = iv.start.astimezone(config.TIMEZONE)
        en = iv.end.astimezone(config.TIMEZONE)
        h = max(iv.duration_hours, 0.0)
        total += h
        per_emp[iv.employee] = per_emp.get(iv.employee, 0.0) + h
        ws.append([st.strftime("%d.%m.%Y"), iv.employee, iv.project or "",
                   st.strftime("%H:%M"), en.strftime("%H:%M"),
                   _fmt_dur(h), round(h, 2), iv.description or "",
                   "manuell" if iv.source == "manual" else "TimeMoto"])

    ws.append([])
    row = ws.max_row + 1
    ws.cell(row, 5, "Summe").font = head
    ws.cell(row, 6, _fmt_dur(total)).font = head
    ws.cell(row, 7, round(total, 2)).font = head
    ws.cell(row, 8, "").font = head

    # Zweites Blatt: Summe je Mitarbeiter
    ws2 = wb.create_sheet("Je Mitarbeiter")
    ws2.append(["Mitarbeiter", "Stunden (Std:Min)", "Stunden (dez.)"])
    for c in ws2[1]:
        c.font = head
        c.fill = fill
    for emp, h in sorted(per_emp.items(), key=lambda kv: kv[1], reverse=True):
        ws2.append([emp, _fmt_dur(h), round(h, 2)])

    for col, w in zip("ABCDEFGHI", (12, 22, 30, 8, 8, 14, 13, 50, 10)):
        ws.column_dimensions[col].width = w
    for col, w in zip("ABC", (24, 16, 14)):
        ws2.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _split_name(full: str) -> str:
    parts = (full or "").split()
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    return f"{parts[-1]}, {' '.join(parts[:-1])}"


def amprion_task(project: str) -> tuple[str, str] | None:
    """Projekt -> (Task-Nr., Label) anhand des Amprion-Mappings, sonst None."""
    p = (project or "").lower()
    for label, nr in config.AMPRION_MAP:
        if label.lower() in p or nr in p:
            return nr, label
    return None


def amprion_xlsx(intervals: list[Interval]) -> bytes:
    """Amprion-'Abgabe'-Format: Datum | (leer) | Task Nr. | (leer) |
    Name (Nachname, Vorname) | Stunden | Tätigkeitsbeschreibung.
    Nur Buchungen auf Amprion-Projektnummern."""
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="EEF5E9")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Abgabe"
    cols = ["Datum", "nicht relevant", "Task Nr.", "nicht relevant",
            "Personen Name", "Stunden", "Taetigkeitbeschreibung"]
    ws.append(cols)
    for c in ws[1]:
        c.font = head
        c.fill = fill

    rows = []
    for iv in intervals:
        m = amprion_task(iv.project)
        if not m:
            continue
        nr, _label = m
        st = iv.start.astimezone(config.TIMEZONE)
        rows.append((st, nr, _split_name(iv.employee),
                     round(max(iv.duration_hours, 0.0), 2), iv.description or ""))
    rows.sort(key=lambda r: (r[0], r[2].lower()))
    for st, nr, name, hours, desc in rows:
        ws.append([st.date(), "", int(nr) if nr.isdigit() else nr, "",
                   name, hours, desc])

    # Legende (Projekt -> Nummer) rechts wie in der Grundlagendatei
    ws.cell(1, 10, "Projekt").font = head
    ws.cell(1, 11, "Task Nr.").font = head
    for i, (label, nr) in enumerate(config.AMPRION_MAP, start=2):
        ws.cell(i, 10, label)
        ws.cell(i, 11, int(nr) if nr.isdigit() else nr)

    for col, w in zip("ABCDEFG", (12, 12, 10, 12, 26, 9, 60)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["J"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
